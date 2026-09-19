"""Report Export Service: Generates styled Excel spreadsheets and PDF analytical reports."""

import io
import uuid
from datetime import datetime, timezone
from typing import Optional

from sqlalchemy.ext.asyncio import AsyncSession

from app.services.analytics_service import analytics_service


class ReportService:
    """Service handling multi-format analytical export generation (Excel .xlsx & PDF)."""

    async def generate_excel_report(
        self,
        db: AsyncSession,
        user_id: uuid.UUID,
        timeframe_days: int = 30,
        platform: Optional[str] = None,
    ) -> io.BytesIO:
        """Generate a multi-tab styled Excel workbook summarizing performance, channels, and content."""
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        # Default sheet
        ws_overview = wb.active
        ws_overview.title = "Overview & Platforms"

        # Styling constants
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # 1. Fetch data
        overview = await analytics_service.get_unified_overview(db, user_id=user_id, platform=platform)
        growth = await analytics_service.get_growth_overview(db, user_id=user_id, platform=platform)
        top_content = await analytics_service.get_top_performing_content(
            db, user_id=user_id, platform=platform, limit=50, days=timeframe_days
        )
        channels = await analytics_service.get_channels_with_metrics(db, user_id=user_id, platform=platform)

        # -------------------------------------------------------------
        # Sheet 1: Overview & KPI Summary
        # -------------------------------------------------------------
        ws_overview["A1"] = "SOCIAL MEDIA ANALYTICS & PERFORMANCE REPORT"
        ws_overview["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
        ws_overview["A2"] = f"Generated at: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | Timeframe: {timeframe_days} Days"
        ws_overview["A2"].font = Font(name="Calibri", size=10, italic=True, color="595959")

        # KPI Summary Cards Table
        ws_overview["A4"] = "METRIC OVERVIEW"
        ws_overview["A4"].font = bold_font
        kpi_headers = ["Metric", "Total Value", "WoW Growth", "MoM Growth"]
        for col_idx, text in enumerate(kpi_headers, 1):
            cell = ws_overview.cell(row=5, column=col_idx, value=text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        kpi_rows = [
            ("Total Channels Connected", overview.total_channels, "-", "-"),
            ("Total Posts Published", overview.total_posts, "-", "-"),
            ("Total Followers / Subscribers", overview.total_followers, f"{growth.followers_wow.growth_rate}%", f"{growth.followers_mom.growth_rate}%"),
            ("Total Views Across Channels", overview.total_views, "-", "-"),
            ("Total Interactions (Likes, Comments, Shares)", overview.total_interactions, "-", "-"),
            ("Average Engagement Rate (ER)", f"{overview.average_engagement_rate}%", "-", "-"),
        ]
        for row_idx, data_tuple in enumerate(kpi_rows, 6):
            for col_idx, val in enumerate(data_tuple, 1):
                cell = ws_overview.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx > 1:
                    cell.alignment = Alignment(horizontal="right")

        # Platform Breakdown Table
        start_plat_row = 14
        ws_overview.cell(row=start_plat_row, column=1, value="PLATFORM PERFORMANCE BREAKDOWN").font = bold_font
        plat_headers = ["Platform", "Active Accounts", "Followers", "Total Views", "Posts", "Interactions", "ER (%)"]
        for col_idx, text in enumerate(plat_headers, 1):
            cell = ws_overview.cell(row=start_plat_row + 1, column=col_idx, value=text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_idx, p in enumerate(overview.platforms, start_plat_row + 2):
            vals = [p.platform.title(), p.accounts_count, p.followers_count, p.views_count, p.posts_count, p.interactions_count, f"{p.engagement_rate}%"]
            for col_idx, v in enumerate(vals, 1):
                c = ws_overview.cell(row=row_idx, column=col_idx, value=v)
                c.border = thin_border
                if col_idx > 1:
                    c.alignment = Alignment(horizontal="right")

        # -------------------------------------------------------------
        # Sheet 2: Channels Detail
        # -------------------------------------------------------------
        ws_chan = wb.create_sheet(title="Connected Channels")
        ws_chan["A1"] = "CONNECTED SOCIAL CHANNELS & GROWTH"
        ws_chan["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")

        chan_headers = ["Channel Name", "Platform", "Platform ID", "Followers", "Total Views", "Posts Count", "WoW Growth", "MoM Growth"]
        for col_idx, h in enumerate(chan_headers, 1):
            cell = ws_chan.cell(row=3, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        growth_map = {cg.channel_id: cg for cg in growth.channels}

        for row_idx, ch in enumerate(channels, 4):
            cg = growth_map.get(ch.id)
            wow = f"{cg.followers_wow.growth_rate}%" if cg else "-"
            mom = f"{cg.followers_mom.growth_rate}%" if cg else "-"
            snap = ch.latest_snapshot
            f_count = snap.followers_count if snap else 0
            v_count = snap.views_count if snap else 0

            row_data = [ch.account_name, ch.platform.title(), ch.platform_account_id, f_count, v_count, ch.posts_count, wow, mom]
            for col_idx, val in enumerate(row_data, 1):
                c = ws_chan.cell(row=row_idx, column=col_idx, value=val)
                c.border = thin_border
                if col_idx in [4, 5, 6, 7, 8]:
                    c.alignment = Alignment(horizontal="right")

        # -------------------------------------------------------------
        # Sheet 3: Top Content Ranking
        # -------------------------------------------------------------
        ws_posts = wb.create_sheet(title="Top Content Performance")
        ws_posts["A1"] = f"TOP PERFORMING POSTS ({timeframe_days} DAYS)"
        ws_posts["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")

        post_headers = ["Title / Excerpt", "Platform", "Channel", "Type", "Published Date", "Views", "Likes", "Comments", "Shares", "ER (%)"]
        for col_idx, h in enumerate(post_headers, 1):
            cell = ws_posts.cell(row=3, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_idx, item in enumerate(top_content.items, 4):
            title = (item.title or item.content or "Untitled")[:60]
            pub_str = item.published_at.strftime("%Y-%m-%d %H:%M")
            row_data = [title, item.platform.title(), item.channel_name, item.post_type.title(), pub_str, item.views_count, item.likes_count, item.comments_count, item.shares_count, f"{item.engagement_rate}%"]
            for col_idx, val in enumerate(row_data, 1):
                c = ws_posts.cell(row=row_idx, column=col_idx, value=val)
                c.border = thin_border
                if col_idx in [6, 7, 8, 9, 10]:
                    c.alignment = Alignment(horizontal="right")

        # Auto-adjust column widths for all sheets
        for sheet in [ws_overview, ws_chan, ws_posts]:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    async def generate_pdf_report(
        self,
        db: AsyncSession,
        user_id: uuid.UUID,
        timeframe_days: int = 30,
        platform: Optional[str] = None,
    ) -> io.BytesIO:
        """Generate a formatted PDF document with tables and key growth metrics."""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        overview = await analytics_service.get_unified_overview(db, user_id=user_id, platform=platform)
        growth = await analytics_service.get_growth_overview(db, user_id=user_id, platform=platform)
        top_content = await analytics_service.get_top_performing_content(
            db, user_id=user_id, platform=platform, limit=10, days=timeframe_days
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            textColor=colors.HexColor("#1F4E78"),
            spaceAfter=6,
        )
        subtitle_style = ParagraphStyle(
            name="SubtitleStyle",
            parent=styles["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#595959"),
            spaceAfter=15,
        )
        section_heading = ParagraphStyle(
            name="SectionHeading",
            parent=styles["Heading2"],
            fontSize=13,
            textColor=colors.HexColor("#1F4E78"),
            spaceBefore=12,
            spaceAfter=8,
        )
        cell_style = ParagraphStyle(
            name="CellStyle",
            parent=styles["Normal"],
            fontSize=9,
            leading=11,
        )
        cell_bold = ParagraphStyle(
            name="CellBold",
            parent=styles["Normal"],
            fontSize=9,
            fontName="Helvetica-Bold",
            leading=11,
        )

        story = []

        # Title & Subtitle
        story.append(Paragraph("SOCIAL MEDIA PERFORMANCE REPORT", title_style))
        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        story.append(Paragraph(f"Generated at: {now_str} | Timeframe: Last {timeframe_days} Days", subtitle_style))

        # KPI Summary Table
        story.append(Paragraph("1. Performance Overview", section_heading))
        kpi_data = [
            [Paragraph("Metric", cell_bold), Paragraph("Value", cell_bold), Paragraph("WoW", cell_bold), Paragraph("MoM", cell_bold)],
            [Paragraph("Total Channels", cell_style), Paragraph(str(overview.total_channels), cell_style), "-", "-"],
            [Paragraph("Total Posts", cell_style), Paragraph(str(overview.total_posts), cell_style), "-", "-"],
            [Paragraph("Total Followers", cell_style), Paragraph(f"{overview.total_followers:,}", cell_style), f"{growth.followers_wow.growth_rate}%", f"{growth.followers_mom.growth_rate}%"],
            [Paragraph("Total Views", cell_style), Paragraph(f"{overview.total_views:,}", cell_style), "-", "-"],
            [Paragraph("Total Interactions", cell_style), Paragraph(f"{overview.total_interactions:,}", cell_style), "-", "-"],
            [Paragraph("Avg Engagement Rate", cell_style), Paragraph(f"{overview.average_engagement_rate}%", cell_style), "-", "-"],
        ]

        t_kpi = Table(kpi_data, colWidths=[200, 140, 100, 100])
        t_kpi.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_kpi)
        story.append(Spacer(1, 15))

        # Platform Breakdown Table
        story.append(Paragraph("2. Platform Breakdown", section_heading))
        plat_data = [
            [Paragraph("Platform", cell_bold), Paragraph("Channels", cell_bold), Paragraph("Followers", cell_bold), Paragraph("Views", cell_bold), Paragraph("ER (%)", cell_bold)]
        ]
        for p in overview.platforms:
            plat_data.append([
                Paragraph(p.platform.title(), cell_style),
                Paragraph(str(p.accounts_count), cell_style),
                Paragraph(f"{p.followers_count:,}", cell_style),
                Paragraph(f"{p.views_count:,}", cell_style),
                Paragraph(f"{p.engagement_rate}%", cell_style),
            ])

        t_plat = Table(plat_data, colWidths=[140, 100, 100, 100, 100])
        t_plat.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5597")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F9")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_plat)
        story.append(Spacer(1, 15))

        # Top Content Table
        if top_content.items:
            story.append(Paragraph("3. Top Performing Content", section_heading))
            post_data = [
                [Paragraph("Title / Content", cell_bold), Paragraph("Platform", cell_bold), Paragraph("Views", cell_bold), Paragraph("Likes", cell_bold), Paragraph("ER (%)", cell_bold)]
            ]
            for item in top_content.items[:8]:
                t = (item.title or item.content or "Untitled")[:40]
                post_data.append([
                    Paragraph(t, cell_style),
                    Paragraph(item.platform.title(), cell_style),
                    Paragraph(f"{item.views_count:,}", cell_style),
                    Paragraph(f"{item.likes_count:,}", cell_style),
                    Paragraph(f"{item.engagement_rate}%", cell_style),
                ])

            t_post = Table(post_data, colWidths=[200, 80, 80, 80, 100])
            t_post.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F9")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t_post)

        doc.build(story)
        output.seek(0)
        return output


report_service = ReportService()
